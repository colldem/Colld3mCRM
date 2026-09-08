from django.utils.translation import gettext as _, gettext_lazy as tr
import mimetypes
import os
import secrets
import uuid
import csv
import io

from django.conf import settings
from django.contrib.auth import get_user_model, update_session_auth_hash
from django.contrib.auth.forms import PasswordChangeForm
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.db import connection, transaction
from django.db.models import Count, Max, Min, Prefetch, Q
from django.http import FileResponse, Http404, HttpResponse, JsonResponse, QueryDict
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.dateparse import parse_date

from .forms import ActivityForm, CompanyForm, DuplicateSettingsForm, PersonForm, ReminderForm, SetupAdminForm, UserProfileForm
from .duplicates import all_company_duplicate_pairs, all_person_duplicate_pairs, find_company_duplicates, find_person_duplicates
from .filters import (
    active_filter_count,
    apply_company_filters,
    apply_contact_filters,
    company_filter_values,
    contact_filter_values,
    filter_chips,
    saved_filter_payload,
)
from .models import Activity, AuditLog, Attachment, Category, Company, CustomField, CustomValue, DuplicateSettings, EmailAddress, Person, PersonCompanyLink, PhoneNumber, PostalAddress, Reminder, SavedFilter, SystemSettings, Tag, Team, UserProfile, WebLink
from .permissions import visible_companies, visible_people, visible_reminders
from .sanitizers import csv_safe, safe_url
from .audit import log as audit_log


def _elided_page_numbers(page):
    """A compact page list for the paginator: ints, with "…" where pages are skipped."""
    raw = page.paginator.get_elided_page_range(page.number, on_each_side=1, on_ends=1)
    return [number if isinstance(number, int) else "…" for number in raw]


def _authorship_context(record, target_type):
    """`Sukūrė` from the record's created_by; `Atnaujino` from the newest audit entry."""
    from .permissions import user_label

    last_edit = (AuditLog.objects.filter(target_type=target_type, target_id=str(record.pk))
                 .exclude(action=AuditLog.CREATE).order_by("-created_at").first())
    updated_by = last_edit.actor_label if last_edit and last_edit.actor_label else user_label(record.created_by)
    return {"created_by_label": user_label(record.created_by), "updated_by_label": updated_by}


def _last_activity_context(last_activity):
    return {"last_activity": last_activity}


def _owner_choices():
    from .permissions import assignable_users

    return assignable_users()


def _owner_label_map(users):
    from .permissions import user_label

    mapping = {str(user.pk): user_label(user) for user in users}
    mapping["none"] = str(tr("Be atsakingo"))
    mapping["me"] = str(tr("Mano įrašai"))
    return mapping


def _can_assign_owner(user):
    from .permissions import has_capability

    return has_capability(user, "can_reassign_owner")


def _list_page_size(request):
    """Rows per list page: an explicit ?page_size wins, else the system default."""
    raw = request.GET.get("page_size")
    if raw in {"25", "50", "100"}:
        return int(raw)
    return SystemSettings.load().default_page_size


def _require_capability(request, capability):
    from .permissions import has_capability

    if not has_capability(request.user, capability):
        raise Http404


def _owner_filter_hint(user, owner_value):
    """Explain an empty result when the user filtered by someone they cannot see."""
    from .permissions import record_visibility, sees_all_records, teammate_ids
    from .models import UserProfile

    if not owner_value or not str(owner_value).isdigit() or sees_all_records(user):
        return ""
    target = int(owner_value)
    vis = record_visibility(user)
    if vis == UserProfile.VISIBILITY_OWN and target != user.pk:
        return str(tr("Jūs matote tik savo įrašus, todėl kito naudotojo įrašai nerodomi."))
    if vis == UserProfile.VISIBILITY_TEAM and target not in teammate_ids(user):
        return str(tr("Jūs matote tik savo komandos įrašus, todėl šio naudotojo įrašai nerodomi."))
    return ""


def _attach_custom_cells(records, custom_columns):
    """Give each record a `.custom_cells` list aligned to `custom_columns` (one query)."""
    for record in records:
        record.custom_cells = []
    if not custom_columns or not records:
        return
    from .custom_fields import _decode, display

    entity_attr = "person" if custom_columns[0].entity == CustomField.PERSON else "company"
    index = {}
    for value in CustomValue.objects.filter(field__in=custom_columns, **{f"{entity_attr}__in": [r.pk for r in records]}).select_related("field"):
        index.setdefault(getattr(value, f"{entity_attr}_id"), {})[value.field_id] = display(value.field, _decode(value.field, value.value))
    for record in records:
        record.custom_cells = [index.get(record.pk, {}).get(field.pk, "") for field in custom_columns]


def _search_results(query, per_group, user=None):
    """Grouped global-search results. `per_group` caps each list; counts are full.

    When `user` is given, results are limited to the records that user may see.
    """
    data = QueryDict(mutable=True)
    data["q"] = query
    terms = query.split()
    text_match = Q()
    for term in terms:
        text_match &= Q(text__icontains=term)

    people = apply_contact_filters(
        visible_people(user, Person.objects.filter(deleted_at__isnull=True)).prefetch_related("emails", "company_links__company"),
        contact_filter_values(data), user,
    ).order_by("last_name", "first_name")
    companies = apply_company_filters(
        visible_companies(user, Company.objects.filter(deleted_at__isnull=True)).prefetch_related("people"),
        company_filter_values(data), user,
    ).order_by("name")
    activities = (
        Activity.objects.filter(deleted_at__isnull=True).filter(text_match)
        .exclude(person__isnull=False, person__deleted_at__isnull=False)
        .exclude(company__isnull=False, company__deleted_at__isnull=False)
        .select_related("person", "company").order_by("-created_at")
    )
    reminders = (
        visible_reminders(user, Reminder.objects.filter(deleted_at__isnull=True).filter(
            Q(person__isnull=True) | Q(person__deleted_at__isnull=True))).filter(text_match)
        .select_related("person", "company").order_by("due_at")
    )
    if user is not None:
        from .permissions import sees_all_records, visible_company_ids, visible_person_ids

        if not sees_all_records(user):
            activities = activities.filter(
                Q(person__pk__in=visible_person_ids(user)) | Q(company__pk__in=visible_company_ids(user))
            )
    return {
        "q": query,
        "people": people[:per_group], "people_count": people.count(),
        "companies": companies[:per_group], "companies_count": companies.count(),
        "activities": activities[:per_group], "activities_count": activities.count(),
        "reminders": reminders[:per_group], "reminders_count": reminders.count(),
    }


@login_required
def global_search(request):
    query = request.GET.get("q", "").strip()
    results = _search_results(query, 50, request.user) if len(query) >= 2 else None
    return render(request, "search.html", {"query": query, "results": results})


@login_required
def search_suggest(request):
    query = request.GET.get("q", "").strip()
    if len(query) < 2:
        return JsonResponse({"q": query, "groups": []})
    results = _search_results(query, 5, request.user)
    search_url = f"{reverse('contacts:search')}?q={query}"
    groups = []
    if results["people_count"]:
        groups.append({"label": str(tr("Kontaktai")), "count": results["people_count"], "url": search_url, "items": [
            {"label": str(person), "sublabel": person.primary_company.name if person.primary_company else "", "url": person.get_absolute_url()}
            for person in results["people"]
        ]})
    if results["companies_count"]:
        groups.append({"label": str(tr("Įmonės")), "count": results["companies_count"], "url": search_url, "items": [
            {"label": company.name, "sublabel": "", "url": company.get_absolute_url()}
            for company in results["companies"]
        ]})
    if results["activities_count"]:
        groups.append({"label": str(tr("Veiklos")), "count": results["activities_count"], "url": search_url, "items": [
            {"label": activity.text[:70], "sublabel": str(activity.person or activity.company or ""),
             "url": (activity.person or activity.company).get_absolute_url() if (activity.person or activity.company) else search_url}
            for activity in results["activities"]
        ]})
    if results["reminders_count"]:
        groups.append({"label": str(tr("Priminimai")), "count": results["reminders_count"], "url": search_url, "items": [
            {"label": reminder.text[:70], "sublabel": str(reminder.record or ""),
             "url": reminder.record_url or reverse("contacts:calendar")}
            for reminder in results["reminders"]
        ]})
    return JsonResponse({"q": query, "groups": groups, "url": search_url})


def health_live(request):
    return JsonResponse({"status": "live"})


def pwa_manifest(request):
    return render(request, "manifest.webmanifest", content_type="application/manifest+json")


def pwa_service_worker(request):
    response = render(request, "sw.js", content_type="application/javascript")
    response["Service-Worker-Allowed"] = "/"
    response["Cache-Control"] = "no-cache"
    return response


def health_ready(request):
    try:
        with connection.cursor() as cursor:
            cursor.execute("SELECT 1")
            cursor.fetchone()
    except Exception:
        return JsonResponse({"status": "not-ready"}, status=503)
    return JsonResponse({"status": "ready"})


def setup_admin(request):
    if get_user_model().objects.exists():
        return redirect("login")
    form = SetupAdminForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        configured_token = settings.CRM_SETUP_TOKEN
        supplied_token = form.cleaned_data["setup_token"]
        if not configured_token or not secrets.compare_digest(configured_token, supplied_token):
            form.add_error("setup_token", "Neteisingas arba neaktyvus diegimo kodas.")
        else:
            with transaction.atomic():
                user = form.save(commit=False)
                user.is_staff = True
                user.is_superuser = True
                user.save()
                UserProfile.objects.update_or_create(user=user, defaults={"role": UserProfile.ROLE_ADMIN})
            return redirect("login")
    return render(request, "registration/setup.html", {"form": form})


@login_required
def contact_list(request):
    redirect_to = _default_filter_redirect(request, "contacts")
    if redirect_to:
        return redirect(redirect_to)
    filter_values = contact_filter_values(request.GET)
    query = filter_values["q"]
    page_size = _list_page_size(request)
    sort_key = request.GET.get("sort", "name")
    direction = "desc" if request.GET.get("direction") == "desc" else "asc"
    sort_map = {"id": ["id"], "name": ["last_name", "first_name"], "company": ["sort_company"], "phone": ["sort_phone"], "email": ["sort_email"], "category": ["sort_category"], "tags": ["sort_tag"], "owner": ["owner__first_name", "owner__last_name", "owner__username"], "last_contact": ["last_contact_at"], "created": ["created_at"], "updated": ["updated_at"]}
    sort_key = sort_key if sort_key in sort_map else "name"
    order_prefix = "-" if direction == "desc" else ""
    order = [f"{order_prefix}{field}" for field in sort_map.get(sort_key, sort_map["name"])]
    people = visible_people(request.user, Person.objects.filter(deleted_at__isnull=True)).select_related("owner").prefetch_related(
        "phones", "emails", "tags", "categories", Prefetch("company_links", queryset=PersonCompanyLink.objects.select_related("company"))
    )
    people = people.annotate(last_contact_at=Max("activities__created_at", filter=Q(activities__deleted_at__isnull=True)))
    people = apply_contact_filters(people, filter_values, request.user)
    people = people.annotate(
        sort_company=Min("company_links__company__name"),
        sort_phone=Min("phones__number"),
        sort_email=Min("emails__email"),
        sort_category=Min("categories__name"),
        sort_tag=Min("tags__name"),
    )
    custom_fields = list(CustomField.objects.filter(entity=CustomField.PERSON))
    allowed_columns = ["company", "phone", "email", "category", "tags", "owner", "last_contact", "updated"] + [field.key for field in custom_fields]
    default_columns = ["company", "phone", "email", "category", "tags", "updated"]
    requested_columns = request.GET.getlist("columns")
    if requested_columns:
        columns = [column for column in requested_columns if column in allowed_columns]
        request.session["contacts_columns"] = columns
    else:
        columns = [column for column in request.session.get("contacts_columns", default_columns) if column in allowed_columns]
    from django.core.paginator import Paginator

    page = Paginator(people.order_by(*order, "id"), page_size).get_page(request.GET.get("page"))
    page.object_list = list(page.object_list)
    custom_columns = [field for field in custom_fields if field.key in columns]
    _attach_custom_cells(page.object_list, custom_columns)
    list_query = request.GET.copy()
    for key in ("page", "sort", "direction"):
        list_query.pop(key, None)
    categories = Category.objects.all()
    tags = Tag.objects.all()
    owner_users = _owner_choices()
    label_maps = {
        "categories": {item.pk: item.name for item in categories},
        "tags": {item.pk: item.name for item in tags},
        "owner": _owner_label_map(owner_users),
        "titles": {"categories": tr("Kategorija"), "tags": tr("Žyma"), **{field.key: field.name for field in custom_fields}},
    }
    return render(request, "contacts/list.html", {
        "page": page, "query": query, "page_size": page_size, "sort": sort_key,
        "page_numbers": _elided_page_numbers(page),
        "direction": direction, "columns": columns, "categories": categories,
        "tags": tags, "filter_values": filter_values,
        "custom_fields": custom_fields, "custom_columns": custom_columns,
        "owner_users": owner_users, "can_assign_owner": _can_assign_owner(request.user),
        "owner_filter_hint": _owner_filter_hint(request.user, filter_values["owner"]),
        "active_filter_count": active_filter_count(filter_values),
        "filter_chips": filter_chips(request.GET, filter_values, label_maps, request.path),
        "saved_filters": SavedFilter.objects.filter(user=request.user, scope="contacts"),
        "list_query": list_query.urlencode(),
    })


def _bulk_add_label(request, queryset, action):
    """Add one tag or category to every record in queryset, keeping the 3-item cap."""
    if action == "add_tag":
        model, field = Tag, "tags"
        item = model.objects.filter(pk=request.POST.get("tag")).first()
        limit_msg = tr("Praleista (jau 3 žymos): %(n)s.")
    else:
        model, field = Category, "categories"
        item = model.objects.filter(pk=request.POST.get("category")).first()
        limit_msg = tr("Praleista (jau 3 kategorijos): %(n)s.")
    if not item:
        return
    added = skipped = 0
    for record in queryset:
        relation = getattr(record, field)
        if relation.filter(pk=item.pk).exists():
            continue
        if relation.count() >= 3:
            skipped += 1
            continue
        try:
            relation.add(item)
            added += 1
            audit_log(AuditLog.UPDATE, request=request, target=record,
                      field=tr("Žyma") if action == "add_tag" else tr("Kategorija"), new=item.name)
        except ValidationError:
            skipped += 1
    if added:
        messages.success(request, tr("Priskirta įrašams: %(n)s.") % {"n": added})
    if skipped:
        messages.error(request, limit_msg % {"n": skipped})


def _bulk_remove_label(request, queryset, action):
    """Remove one tag or category from every record in queryset."""
    if action == "remove_tag":
        model, field, key = Tag, "tags", "tag"
    else:
        model, field, key = Category, "categories", "category"
    item = model.objects.filter(pk=request.POST.get(key)).first()
    if not item:
        return
    removed = 0
    for record in queryset:
        relation = getattr(record, field)
        if relation.filter(pk=item.pk).exists():
            relation.remove(item)
            removed += 1
            audit_log(AuditLog.UPDATE, request=request, target=record,
                      field=tr("Žyma") if action == "remove_tag" else tr("Kategorija"), old=item.name)
    if removed:
        messages.success(request, tr("Nuimta nuo įrašų: %(n)s.") % {"n": removed})


def _bulk_assign_owner(request, queryset):
    """Set (or clear, when value is empty) the owner for every record in queryset."""
    if not _can_assign_owner(request.user):
        messages.error(request, tr("Neturite teisės keisti atsakingo naudotojo."))
        return
    raw = request.POST.get("owner", "").strip()
    owner = None
    if raw:
        owner = get_user_model().objects.filter(pk=raw, is_active=True).first()
        if not owner:
            messages.error(request, tr("Pasirinktas naudotojas nerastas."))
            return
    from .permissions import user_label

    records = list(queryset.select_related("owner"))
    for record in records:
        if record.owner_id != (owner.pk if owner else None):
            audit_log(AuditLog.UPDATE, request=request, target=record, field=tr("Atsakingas"),
                      old=user_label(record.owner), new=user_label(owner))
    count = queryset.update(owner=owner, updated_at=timezone.now())
    if owner:
        messages.success(request, tr("Atsakingas priskirtas įrašams: %(n)s.") % {"n": count})
    else:
        messages.success(request, tr("Atsakingas pašalintas nuo įrašų: %(n)s.") % {"n": count})


@login_required
def contact_bulk_action(request):
    if request.method != "POST":
        return redirect("contacts:list")
    action = request.POST.get("action")
    ids = request.POST.getlist("selected")
    if not ids:
        return redirect("contacts:list")
    people = visible_people(request.user, Person.objects.filter(pk__in=ids, deleted_at__isnull=True))
    if action == "archive":
        _require_capability(request, "can_delete")
        for person in people:
            audit_log(AuditLog.ARCHIVE, request=request, target=person)
        count = people.update(deleted_at=timezone.now())
        messages.success(request, tr("Archyvuota kontaktų: %(count)s.") % {"count": count})
    elif action in {"add_tag", "add_category"}:
        _require_capability(request, "can_bulk_edit")
        _bulk_add_label(request, people, action)
    elif action in {"remove_tag", "remove_category"}:
        _require_capability(request, "can_bulk_edit")
        _bulk_remove_label(request, people, action)
    elif action == "assign_owner":
        _bulk_assign_owner(request, people)
    return redirect("contacts:list")


@login_required
def company_bulk_action(request):
    if request.method != "POST":
        return redirect("contacts:company-list")
    action = request.POST.get("action")
    ids = request.POST.getlist("selected")
    if not ids:
        return redirect("contacts:company-list")
    companies = visible_companies(request.user, Company.objects.filter(pk__in=ids, deleted_at__isnull=True))
    if action == "archive":
        _require_capability(request, "can_delete")
        for company in companies:
            audit_log(AuditLog.ARCHIVE, request=request, target=company)
        count = companies.update(deleted_at=timezone.now())
        if count:
            messages.success(request, tr("Archyvuota įmonių: %(count)s.") % {"count": count})
    elif action in {"add_tag", "add_category"}:
        _require_capability(request, "can_bulk_edit")
        _bulk_add_label(request, companies, action)
    elif action in {"remove_tag", "remove_category"}:
        _require_capability(request, "can_bulk_edit")
        _bulk_remove_label(request, companies, action)
    elif action == "assign_owner":
        _bulk_assign_owner(request, companies)
    return redirect("contacts:company-list")


@login_required
def contact_archive(request, pk):
    _require_capability(request, "can_delete")
    person = get_object_or_404(visible_people(request.user), pk=pk)
    if request.method == "POST" and person.deleted_at is None:
        person.deleted_at = timezone.now()
        person.save(update_fields=["deleted_at", "updated_at"])
        audit_log(AuditLog.ARCHIVE, request=request, target=person)
        messages.success(request, tr("Kontaktas perkeltas į archyvą."))
    return redirect("contacts:list")


@login_required
def company_archive(request, pk):
    _require_capability(request, "can_delete")
    company = get_object_or_404(visible_companies(request.user), pk=pk)
    if request.method == "POST" and company.deleted_at is None:
        company.deleted_at = timezone.now()
        company.save(update_fields=["deleted_at", "updated_at"])
        audit_log(AuditLog.ARCHIVE, request=request, target=company)
        messages.success(request, tr("Įmonė perkelta į archyvą."))
    return redirect("contacts:company-list")


@login_required
def archive_list(request):
    people = visible_people(request.user, Person.objects.filter(deleted_at__isnull=False, merged_into__isnull=True)).order_by("-deleted_at")
    companies = visible_companies(request.user, Company.objects.filter(deleted_at__isnull=False, merged_into__isnull=True)).order_by("-deleted_at")
    return render(request, "archive.html", {"people": people, "companies": companies})


@login_required
def contact_restore(request, pk):
    _require_capability(request, "can_delete")
    person = get_object_or_404(visible_people(request.user), pk=pk, merged_into__isnull=True)
    if request.method == "POST" and person.deleted_at is not None:
        person.deleted_at = None
        person.save(update_fields=["deleted_at", "updated_at"])
        audit_log(AuditLog.RESTORE, request=request, target=person)
        messages.success(request, tr("Kontaktas atkurtas."))
    return redirect("contacts:archive-list")


@login_required
def company_restore(request, pk):
    _require_capability(request, "can_delete")
    company = get_object_or_404(visible_companies(request.user), pk=pk, merged_into__isnull=True)
    if request.method == "POST" and company.deleted_at is not None:
        company.deleted_at = None
        company.save(update_fields=["deleted_at", "updated_at"])
        audit_log(AuditLog.RESTORE, request=request, target=company)
        messages.success(request, tr("Įmonė atkurta."))
    return redirect("contacts:archive-list")


def _default_filter_redirect(request, scope):
    """Apply the user's default saved filter once per session on the first
    parameter-free visit to the list. Returns a URL to redirect to, or None."""
    seen_key = f"{scope}_list_seen"
    first_visit = not request.session.get(seen_key)
    request.session[seen_key] = True
    if not first_visit or request.GET:
        return None
    default = SavedFilter.objects.filter(user=request.user, scope=scope, is_default=True).first()
    if default and default.query_string:
        return f"{request.path}?{default.query_string}"
    return None


@login_required
def saved_filter_create(request):
    if request.method == "POST":
        name = request.POST.get("name", "").strip()
        filters = saved_filter_payload(request.POST, "contacts")
        if name:
            SavedFilter.objects.update_or_create(user=request.user, scope="contacts", name=name, defaults={"filters": filters})
            messages.success(request, tr("Filtras išsaugotas."))
    return redirect("contacts:list")


@login_required
def company_saved_filter_create(request):
    if request.method == "POST":
        name = request.POST.get("name", "").strip()
        filters = saved_filter_payload(request.POST, "companies")
        if name:
            SavedFilter.objects.update_or_create(user=request.user, scope="companies", name=name, defaults={"filters": filters})
            messages.success(request, tr("Įmonių sąrašas išsaugotas."))
    return redirect("contacts:company-list")


@login_required
def saved_filter_update(request, pk):
    saved = get_object_or_404(SavedFilter, pk=pk, user=request.user)
    target = "contacts:company-list" if saved.scope == "companies" else "contacts:list"
    if request.method == "POST":
        action = request.POST.get("action")
        if action == "delete":
            saved.delete()
            messages.success(request, tr("Filtras pašalintas."))
        elif action == "rename":
            name = request.POST.get("name", "").strip()
            if not name:
                messages.error(request, tr("Filtro pavadinimas negali būti tuščias."))
            elif SavedFilter.objects.filter(user=request.user, scope=saved.scope, name=name).exclude(pk=saved.pk).exists():
                messages.error(request, tr("Toks filtro pavadinimas jau yra."))
            else:
                saved.name = name
                saved.save(update_fields=["name"])
                messages.success(request, tr("Filtras pervadintas."))
        elif action == "default":
            make_default = not saved.is_default
            SavedFilter.objects.filter(user=request.user, scope=saved.scope).exclude(pk=saved.pk).update(is_default=False)
            SavedFilter.objects.filter(pk=saved.pk).update(is_default=make_default)
            messages.success(request, tr("Numatytasis filtras nustatytas.") if make_default else tr("Numatytasis filtras išjungtas."))
    return redirect(target)


@login_required
def settings_page(request):
    if request.method == "POST" and request.POST.get("kind") in {"tag", "category"}:
        return settings_taxonomy(request, request.POST["kind"])
    form = UserProfileForm(request.POST or None, request.FILES or None, user=request.user)
    if request.method == "POST" and form.is_valid():
        changed = list(form.changed_data)
        profile = form.save()
        if changed:
            audit_log(AuditLog.SETTING, request=request, target_type="setting",
                      target_label=str(tr("Profilis")), new=", ".join(changed))
        messages.success(request, tr("Profilis atnaujintas."))
        response = redirect("contacts:settings")
        response.set_cookie(settings.LANGUAGE_COOKIE_NAME, profile.language, max_age=365 * 24 * 60 * 60, samesite="Lax")
        return response
    return render(request, "settings/profile.html", {
        "form": form,
        "password_form": PasswordChangeForm(request.user),
        "settings_section": "profile",
    })


@login_required
def settings_custom_fields(request):
    from .models import CustomField

    _require_capability(request, "can_manage_custom_fields")
    if request.method == "POST":
        name = request.POST.get("name", "").strip()[:60]
        field_type = request.POST.get("field_type", CustomField.TEXT)
        entity = request.POST.get("entity", "both")
        options = [line.strip() for line in request.POST.get("options", "").splitlines() if line.strip()][:20]
        valid_type = field_type in dict(CustomField.TYPE_CHOICES)
        needs_options = field_type in (CustomField.SELECT, CustomField.MULTISELECT)
        targets = [CustomField.PERSON, CustomField.COMPANY] if entity == "both" else [entity] if entity in (CustomField.PERSON, CustomField.COMPANY) else []
        if not name or not valid_type or not targets or (needs_options and not options):
            messages.error(request, tr("Nurodykite pavadinimą, tipą, paskirtį ir (jei reikia) reikšmes."))
        else:
            created = 0
            for target in targets:
                if not CustomField.objects.filter(entity=target, name__iexact=name).exists():
                    order = (CustomField.objects.filter(entity=target).count())
                    field = CustomField.objects.create(entity=target, name=name, field_type=field_type, options=options if needs_options else [], order=order)
                    created += 1
                    audit_log(AuditLog.CREATE, request=request, target_type="custom_field", target_id=field.pk,
                              target_label=f"{name} ({field.get_entity_display()})", new=field.get_field_type_display())
            messages.success(request, tr("Laukas pridėtas.") if created else tr("Toks laukas jau yra."))
        return redirect("contacts:settings-custom-fields")
    people_fields = CustomField.objects.filter(entity=CustomField.PERSON)
    company_fields = CustomField.objects.filter(entity=CustomField.COMPANY)
    return render(request, "settings/custom_fields.html", {
        "settings_section": "custom-fields",
        "people_fields": people_fields, "company_fields": company_fields,
        "type_choices": CustomField.TYPE_CHOICES,
    })


@login_required
def custom_field_delete(request, pk):
    from .models import CustomField

    _require_capability(request, "can_manage_custom_fields")
    field = get_object_or_404(CustomField, pk=pk)
    if request.method == "POST":
        label = f"{field.name} ({field.get_entity_display()})"
        field.delete()
        audit_log(AuditLog.DELETE, request=request, target_type="custom_field", target_id=pk, target_label=label)
        messages.success(request, tr("Laukas pašalintas."))
    return redirect("contacts:settings-custom-fields")


def _create_crm_user(request):
    from django.contrib.auth.password_validation import validate_password

    User = get_user_model()
    username = request.POST.get("username", "").strip()
    role = request.POST.get("role", UserProfile.ROLE_MEMBER)
    password = request.POST.get("password", "")
    if not username or User.objects.filter(username__iexact=username).exists():
        messages.error(request, tr("Toks naudotojo vardas jau naudojamas."))
        return
    if role not in dict(UserProfile.ROLE_CHOICES):
        role = UserProfile.ROLE_MEMBER
    try:
        validate_password(password)
    except ValidationError as error:
        messages.error(request, " ".join(error.messages))
        return
    visibility = request.POST.get("record_visibility")
    if visibility not in dict(UserProfile.VISIBILITY_CHOICES):
        visibility = UserProfile.VISIBILITY_OWN if role == UserProfile.ROLE_RESTRICTED else UserProfile.VISIBILITY_ALL
    user = User.objects.create_user(
        username=username, email=request.POST.get("email", "").strip(), password=password,
        first_name=request.POST.get("first_name", "").strip(), last_name=request.POST.get("last_name", "").strip(),
        is_staff=role == UserProfile.ROLE_ADMIN,
    )
    UserProfile.objects.update_or_create(user=user, defaults={"role": role, "record_visibility": visibility})
    audit_log(AuditLog.CREATE, request=request, target=user, target_type="user", field=str(tr("Rolė")),
              new=dict(UserProfile.ROLE_CHOICES).get(role, role))
    messages.success(request, tr("Naudotojas sukurtas."))


def _update_crm_user(request, target):
    from .permissions import active_admin_ids, is_admin, role_of

    role = request.POST.get("role", UserProfile.ROLE_MEMBER)
    if role not in dict(UserProfile.ROLE_CHOICES):
        role = UserProfile.ROLE_MEMBER
    active = request.POST.get("active") == "1"
    admins = active_admin_ids()
    losing_admin = is_admin(target) and (role != UserProfile.ROLE_ADMIN or not active)
    if losing_admin and admins <= {target.pk}:
        messages.error(request, tr("Turi likti bent vienas aktyvus administratorius."))
        return
    if target.is_superuser and (role != UserProfile.ROLE_ADMIN or not active):
        messages.error(request, tr("Pagrindinio administratoriaus keisti negalima."))
        return
    old_role = role_of(target)
    old_active = target.is_active
    old_profile = UserProfile.objects.filter(user=target).first()
    old_visibility = old_profile.record_visibility if old_profile else UserProfile.VISIBILITY_ALL
    visibility = request.POST.get("record_visibility")
    if visibility not in dict(UserProfile.VISIBILITY_CHOICES):
        visibility = old_visibility
    target.is_active = active
    target.is_staff = role == UserProfile.ROLE_ADMIN
    target.save(update_fields=["is_active", "is_staff"])
    UserProfile.objects.update_or_create(user=target, defaults={"role": role, "record_visibility": visibility})
    labels = dict(UserProfile.ROLE_CHOICES)
    vis_labels = dict(UserProfile.VISIBILITY_CHOICES)
    if old_role != role:
        audit_log(AuditLog.UPDATE, request=request, target=target, target_type="user", field=str(tr("Rolė")),
                  old=labels.get(old_role, old_role), new=labels.get(role, role))
    if old_visibility != visibility:
        audit_log(AuditLog.UPDATE, request=request, target=target, target_type="user", field=str(tr("Matomumas")),
                  old=vis_labels.get(old_visibility, old_visibility), new=vis_labels.get(visibility, visibility))
    if old_active != active:
        audit_log(AuditLog.UPDATE, request=request, target=target, target_type="user", field=str(tr("Būsena")),
                  old=str(tr("aktyvus")) if old_active else str(tr("išjungtas")),
                  new=str(tr("aktyvus")) if active else str(tr("išjungtas")))
    messages.success(request, tr("Naudotojas atnaujintas."))


def _reset_crm_user_password(request, target):
    from django.contrib.auth.password_validation import validate_password

    if target.is_superuser and not request.user.is_superuser:
        messages.error(request, tr("Pagrindinio administratoriaus slaptažodžio keisti negalima."))
        return
    password = request.POST.get("password", "")
    try:
        validate_password(password, user=target)
    except ValidationError as error:
        messages.error(request, " ".join(error.messages))
        return
    target.set_password(password)
    target.save(update_fields=["password"])
    if target.pk == request.user.pk:
        update_session_auth_hash(request, target)
    audit_log(AuditLog.UPDATE, request=request, target=target, target_type="user", field=str(tr("Slaptažodis")),
              new=str(tr("atstatytas")))
    messages.success(request, tr("Slaptažodis atstatytas."))


@login_required
def settings_users(request):
    from .permissions import is_admin, role_of

    if not is_admin(request.user):
        raise Http404
    User = get_user_model()
    if request.method == "POST":
        action = request.POST.get("action")
        if action == "create":
            _create_crm_user(request)
        elif action in ("update", "reset"):
            target = User.objects.filter(pk=request.POST.get("user_id")).first()
            if target and action == "update":
                _update_crm_user(request, target)
            elif target:
                _reset_crm_user_password(request, target)
        return redirect("contacts:settings-users")
    from .permissions import record_visibility as effective_visibility
    rows = [{
        "user": user,
        "name": user.get_full_name().strip() or user.get_username(),
        "role": role_of(user),
        "visibility": getattr(getattr(user, "crm_profile", None), "record_visibility", UserProfile.VISIBILITY_ALL),
        "effective_visibility": effective_visibility(user),
        "teams": ", ".join(user.crm_teams.values_list("name", flat=True)),
        "is_self": user.pk == request.user.pk,
        "protected": user.is_superuser,
    } for user in User.objects.select_related("crm_profile").prefetch_related("crm_teams").order_by("username")]
    return render(request, "settings/users.html", {
        "settings_section": "users", "rows": rows, "role_choices": UserProfile.ROLE_CHOICES,
        "visibility_choices": UserProfile.VISIBILITY_CHOICES,
    })


@login_required
def settings_teams(request):
    from .permissions import is_admin

    if not is_admin(request.user):
        raise Http404
    User = get_user_model()
    if request.method == "POST":
        action = request.POST.get("action")
        if action == "create":
            name = request.POST.get("name", "").strip()[:80]
            if not name:
                messages.error(request, tr("Įveskite komandos pavadinimą."))
            elif Team.objects.filter(name__iexact=name).exists():
                messages.error(request, tr("Tokia komanda jau yra."))
            else:
                team = Team.objects.create(name=name)
                audit_log(AuditLog.CREATE, request=request, target=team, target_type="team")
                messages.success(request, tr("Komanda sukurta."))
        else:
            team = Team.objects.filter(pk=request.POST.get("team_id")).first()
            if not team:
                pass
            elif action == "delete":
                audit_log(AuditLog.DELETE, request=request, target_type="team", target_id=team.pk, target_label=team.name)
                team.delete()
                messages.success(request, tr("Komanda pašalinta."))
            elif action == "update":
                name = request.POST.get("name", "").strip()[:80]
                visibility = request.POST.get("visibility") if request.POST.get("visibility") in dict(Team.VISIBILITY_CHOICES) else team.visibility
                member_ids = [int(v) for v in request.POST.getlist("members") if v.isdigit()]
                new_members = list(User.objects.filter(pk__in=member_ids, is_active=True))
                old = {"name": team.name, "visibility": team.visibility, "members": sorted(team.members.values_list("username", flat=True))}
                if name and name.lower() != team.name.lower() and Team.objects.filter(name__iexact=name).exclude(pk=team.pk).exists():
                    messages.error(request, tr("Tokia komanda jau yra."))
                    return redirect("contacts:settings-teams")
                if name:
                    team.name = name
                team.visibility = visibility
                team.save(update_fields=["name", "visibility"])
                team.members.set(new_members)
                new = {"name": team.name, "visibility": team.visibility, "members": sorted(m.username for m in new_members)}
                for key in ("name", "visibility", "members"):
                    if old[key] != new[key]:
                        audit_log(AuditLog.UPDATE, request=request, target=team, target_type="team", field=key,
                                  old=", ".join(old[key]) if isinstance(old[key], list) else old[key],
                                  new=", ".join(new[key]) if isinstance(new[key], list) else new[key])
                messages.success(request, tr("Komanda atnaujinta."))
        return redirect("contacts:settings-teams")
    users = list(User.objects.filter(is_active=True).order_by("first_name", "last_name", "username"))
    teams = Team.objects.prefetch_related("members").all()
    rows = [{"team": team, "member_ids": set(team.members.values_list("id", flat=True))} for team in teams]
    return render(request, "settings/teams.html", {
        "settings_section": "teams", "rows": rows, "users": users,
        "visibility_choices": Team.VISIBILITY_CHOICES,
    })


@login_required
def settings_permissions(request):
    from .permissions import CAPABILITIES, capability_matrix, is_admin
    from .models import RolePermissions

    if not is_admin(request.user):
        raise Http404
    editable_roles = [UserProfile.ROLE_MEMBER, UserProfile.ROLE_RESTRICTED]
    if request.method == "POST":
        for role in editable_roles:
            granted = set(request.POST.getlist("cap_" + role))
            permissions = {key: (key in granted) for key, _label in CAPABILITIES}
            RolePermissions.objects.update_or_create(role=role, defaults={"permissions": permissions})
        audit_log(AuditLog.SETTING, request=request, target_type="setting",
                  target_label=str(tr("Rolės ir teisės")), new=str(tr("atnaujinta")))
        messages.success(request, tr("Teisės išsaugotos."))
        return redirect("contacts:settings-permissions")
    matrix = capability_matrix()
    role_labels = dict(UserProfile.ROLE_CHOICES)
    return render(request, "settings/permissions.html", {
        "settings_section": "permissions",
        "capabilities": CAPABILITIES,
        "roles": [{"key": role, "label": role_labels.get(role, role), "caps": matrix[role]} for role in editable_roles],
    })


@login_required
def settings_audit(request):
    from .permissions import has_capability

    if not has_capability(request.user, "can_view_audit"):
        raise Http404
    entries = AuditLog.objects.select_related("actor").all()
    actor_id = request.GET.get("actor", "").strip()
    action = request.GET.get("action", "").strip()
    date_from = request.GET.get("date_from", "").strip()
    date_to = request.GET.get("date_to", "").strip()
    if actor_id.isdigit():
        entries = entries.filter(actor_id=actor_id)
    if action in dict(AuditLog.ACTION_CHOICES):
        entries = entries.filter(action=action)
    if date_from and parse_date(date_from):
        entries = entries.filter(created_at__date__gte=date_from)
    if date_to and parse_date(date_to):
        entries = entries.filter(created_at__date__lte=date_to)
    from django.core.paginator import Paginator

    page = Paginator(entries, 100).get_page(request.GET.get("page"))
    list_query = request.GET.copy()
    list_query.pop("page", None)
    return render(request, "settings/audit.html", {
        "settings_section": "audit",
        "page": page,
        "page_numbers": _elided_page_numbers(page),
        "actors": get_user_model().objects.filter(audit_entries__isnull=False).distinct().order_by("username"),
        "action_choices": AuditLog.ACTION_CHOICES,
        "filter_actor": actor_id, "filter_action": action,
        "filter_date_from": date_from, "filter_date_to": date_to,
        "list_query": list_query.urlencode(),
    })


@login_required
def settings_password(request):
    # The password form lives on the profile page; this endpoint only takes its POST.
    if request.method != "POST":
        return redirect("contacts:settings")
    form = PasswordChangeForm(request.user, request.POST)
    if form.is_valid():
        form.save()
        update_session_auth_hash(request, form.user)
        audit_log(AuditLog.UPDATE, request=request, target=request.user, target_type="user",
                  field=str(tr("Slaptažodis")), new=str(tr("pakeistas")))
        messages.success(request, tr("Slaptažodis pakeistas."))
        return redirect("contacts:settings")
    return render(request, "settings/profile.html", {
        "form": UserProfileForm(user=request.user),
        "password_form": form,
        "settings_section": "profile",
    })


_EXPORT_README = (
    "CRM duomenų eksportas / CRM data export\n"
    "======================================\n\n"
    "data.json  - visų CRM duomenų kopija (Django dumpdata).\n"
    "media/     - įrašų priedai ir profilių nuotraukos.\n\n"
    "Atkūrimas / restore:\n"
    "  1. python manage.py migrate\n"
    "  2. python manage.py loaddata data.json\n"
    "  3. media/ turinį nukopijuoti į runtime/media/\n"
)


@login_required
def settings_data_export(request):
    if not request.user.is_staff:
        raise Http404
    if request.method == "POST":
        import io
        import zipfile
        from django.core.management import call_command

        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            dump = io.StringIO()
            call_command(
                "dumpdata", "auth.Group", "auth.User", "contacts",
                natural_foreign=True, natural_primary=True, indent=2, stdout=dump,
            )
            archive.writestr("data.json", dump.getvalue())
            archive.writestr("README.txt", _EXPORT_README)
            media_root = str(settings.MEDIA_ROOT)
            if os.path.isdir(media_root):
                for folder, _subdirs, names in os.walk(media_root):
                    for name in names:
                        full = os.path.join(folder, name)
                        archive.write(full, os.path.join("media", os.path.relpath(full, media_root)))
        response = HttpResponse(buffer.getvalue(), content_type="application/zip")
        response["Content-Disposition"] = 'attachment; filename="crm-eksportas-%s.zip"' % timezone.now().strftime("%Y%m%d-%H%M")
        audit_log(AuditLog.EXPORT, request=request, target_type="export", target_label=str(tr("Pilna atsarginė kopija (ZIP)")))
        return response
    return render(request, "settings/data_export.html", {"settings_section": "data-export"})


@login_required
def profile_avatar(request):
    profile = UserProfile.objects.filter(user=request.user).first()
    if not profile or not profile.avatar:
        raise Http404
    try:
        response = FileResponse(
            profile.avatar.open("rb"),
            content_type=mimetypes.guess_type(profile.avatar.name)[0] or "application/octet-stream",
        )
    except FileNotFoundError:
        raise Http404("Profilio nuotrauka nerasta.")
    response["Cache-Control"] = "private, no-store"
    return response


@login_required
def settings_taxonomy(request, kind):
    _require_capability(request, "can_manage_taxonomy")
    model = Tag if kind == "tag" else Category if kind == "category" else None
    if not model:
        raise Http404
    if request.method == "POST":
        name = request.POST.get("name", "").strip()
        item_id = request.POST.get("item_id", "").strip()
        if request.POST.get("delete") and item_id:
            item = get_object_or_404(model, pk=item_id)
            label = item.name
            item.delete()
            audit_log(AuditLog.SETTING, request=request, target_type="setting",
                      target_label=(str(tr("Žyma")) if kind == "tag" else str(tr("Kategorija"))) + f": {label}", new=str(tr("ištrinta")))
            messages.success(request, tr("Žyma ištrinta.") if kind == "tag" else tr("Kategorija ištrinta."))
            return redirect("contacts:settings-tags" if kind == "tag" else "contacts:settings-categories")
        if not name:
            messages.error(request, tr("Žyma negali būti tuščia.") if kind == "tag" else tr("Kategorija negali būti tuščia."))
        elif model.objects.filter(name__iexact=name).exclude(pk=item_id or None).exists():
            messages.error(request, tr("Tokia žyma jau yra.") if kind == "tag" else tr("Tokia kategorija jau yra."))
        elif item_id:
            item = get_object_or_404(model, pk=item_id)
            item.name = name
            update_fields = ["name"]
            if kind == "tag":
                color = request.POST.get("color", "")
                if color and color not in dict(Tag.COLOR_CHOICES):
                    messages.error(request, tr("Pasirinkta netinkama žymos spalva."))
                    return redirect("contacts:settings-tags")
                item.color = color
                update_fields.append("color")
            item.save(update_fields=update_fields)
            audit_log(AuditLog.SETTING, request=request, target_type="setting",
                      target_label=(str(tr("Žyma")) if kind == "tag" else str(tr("Kategorija"))) + f": {name}", new=str(tr("atnaujinta")))
            messages.success(request, tr("Žyma atnaujinta.") if kind == "tag" else tr("Kategorija atnaujinta."))
        else:
            item = model.objects.create(name=name)
            if kind == "tag" and request.POST.get("color") in dict(Tag.COLOR_CHOICES):
                item.color = request.POST["color"]
                item.save(update_fields=["color"])
            audit_log(AuditLog.SETTING, request=request, target_type="setting",
                      target_label=(str(tr("Žyma")) if kind == "tag" else str(tr("Kategorija"))) + f": {name}", new=str(tr("sukurta")))
            messages.success(request, tr("Žyma pridėta.") if kind == "tag" else tr("Kategorija pridėta."))
        return redirect("contacts:settings-tags" if kind == "tag" else "contacts:settings-categories")
    items = model.objects.annotate(people_count=Count("people", distinct=True), companies_count=Count("companies", distinct=True))
    return render(request, "settings/taxonomy.html", {"items": items, "kind": kind, "tag_colors": Tag.COLOR_CHOICES, "settings_section": "tags" if kind == "tag" else "categories"})


@login_required
def settings_duplicates(request):
    from .permissions import is_admin

    if not is_admin(request.user):
        raise Http404
    duplicate_settings = DuplicateSettings.load()
    form = DuplicateSettingsForm(request.POST or None, instance=duplicate_settings)
    if request.method == "POST" and form.is_valid():
        changed = list(form.changed_data)
        form.save()
        if changed:
            audit_log(AuditLog.SETTING, request=request, target_type="setting",
                      target_label=str(tr("Dublikatų tikrinimas")), new=", ".join(changed))
        messages.success(request, tr("Dublikatų tikrinimo nustatymai išsaugoti."))
        return redirect("contacts:settings-duplicates")
    return render(request, "settings/duplicates.html", {"form": form, "settings_section": "duplicates"})


@login_required
def settings_system(request):
    from .forms import SystemSettingsForm
    from .permissions import is_admin

    if not is_admin(request.user):
        raise Http404
    system = SystemSettings.load()
    form = SystemSettingsForm(request.POST or None, instance=system)
    if request.method == "POST" and form.is_valid():
        changed = list(form.changed_data)
        form.save()
        if changed:
            audit_log(AuditLog.SETTING, request=request, target_type="setting",
                      target_label=str(tr("Sistemos nustatymai")), new=", ".join(changed))
        messages.success(request, tr("Sistemos nustatymai išsaugoti."))
        return redirect("contacts:settings-system")
    return render(request, "settings/system.html", {"form": form, "settings_section": "system"})


@login_required
def settings_import(request):
    from .forms import ImportSettingsForm
    from .permissions import is_admin

    if not is_admin(request.user):
        raise Http404
    form = ImportSettingsForm(request.POST or None, instance=SystemSettings.load())
    if request.method == "POST" and form.is_valid():
        changed = list(form.changed_data)
        form.save()
        if changed:
            audit_log(AuditLog.SETTING, request=request, target_type="setting",
                      target_label=str(tr("Importo nustatymai")), new=", ".join(changed))
        messages.success(request, tr("Importo nustatymai išsaugoti."))
        return redirect("contacts:settings-import")
    return render(request, "settings/import_settings.html", {"form": form, "settings_section": "import"})


@login_required
def documentation_page(request):
    topics = (
        ("overview", tr("Pradžia")),
        ("installation", tr("Diegimas")),
        ("screens", tr("Langai ir mygtukai")),
        ("user", tr("Naudotojui")),
        ("admin", tr("Administratoriui")),
        ("data", tr("Duomenys ir sauga")),
        ("backup", tr("Kopijos ir atkūrimas")),
    )
    allowed_topics = {key for key, _label in topics}
    topic = request.GET.get("topic", "overview")
    if topic not in allowed_topics:
        topic = "overview"
    try:
        crm_version = (settings.BASE_DIR / "VERSION").read_text().strip()
    except OSError:
        crm_version = ""
    return render(request, "settings/documentation.html", {
        "settings_section": "documentation",
        "documentation_topic": topic,
        "documentation_topics": topics,
        "crm_version": crm_version,
    })


@login_required
def duplicate_list(request):
    duplicate_settings = DuplicateSettings.load()
    pairs = (all_person_duplicate_pairs(duplicate_settings.level) + all_company_duplicate_pairs(duplicate_settings.level)) if duplicate_settings.enabled else []
    from .permissions import can_see_company, can_see_person

    def _visible_pair(pair):
        left, right = pair["left"], pair["right"]
        checker = can_see_person if isinstance(left, Person) else can_see_company
        return checker(request.user, left) and checker(request.user, right)

    pairs = [pair for pair in pairs if _visible_pair(pair)]
    return render(request, "duplicates/list.html", {"pairs": pairs, "duplicate_settings": duplicate_settings})


@login_required
def duplicate_merge_all(request):
    if request.method != "POST":
        return redirect("contacts:duplicate-list")
    _require_capability(request, "can_merge_duplicates")
    duplicate_settings = DuplicateSettings.load()
    if not duplicate_settings.enabled:
        messages.error(request, tr("Dublikatų tikrinimas išjungtas."))
        return redirect("contacts:duplicate-list")

    from .merging import merge_companies, merge_people
    from .permissions import can_see_company, can_see_person

    merged = 0
    plans = (
        (all_person_duplicate_pairs(duplicate_settings.level), can_see_person, merge_people),
        (all_company_duplicate_pairs(duplicate_settings.level), can_see_company, merge_companies),
    )
    for pairs, checker, merge_fn in plans:
        for pair in pairs:
            keep, drop = pair["left"], pair["right"]
            keep.refresh_from_db()
            drop.refresh_from_db()
            # A record can already be gone via an earlier merge in a duplicate chain.
            if keep.deleted_at or drop.deleted_at or keep.merged_into_id or drop.merged_into_id:
                continue
            if not (checker(request.user, keep) and checker(request.user, drop)):
                continue
            try:
                target = merge_fn(drop.pk, keep.pk)
            except ValidationError:
                continue
            audit_log(AuditLog.MERGE, request=request, target=target, old=str(drop), new=str(target),
                      detail={"source_id": drop.pk, "target_id": keep.pk, "bulk": True})
            merged += 1
    if merged:
        messages.success(request, tr("Sujungta dublikatų porų: %(n)s.") % {"n": merged})
    else:
        messages.info(request, tr("Sujungiamų dublikatų nerasta."))
    return redirect("contacts:duplicate-list")


@login_required
def duplicate_merge(request, kind, source_pk, target_pk):
    if request.method != "POST":
        return HttpResponse(status=405)
    _require_capability(request, "can_merge_duplicates")
    if kind not in {"person", "company"} or source_pk == target_pk:
        return HttpResponse(tr("Netinkami sujungimo duomenys."), status=400)
    duplicate_settings = DuplicateSettings.load()
    if not duplicate_settings.enabled:
        return HttpResponse(tr("Dublikatų tikrinimas išjungtas."), status=400)

    model = Person if kind == "person" else Company
    source = model.objects.filter(pk=source_pk).first()
    target = model.objects.filter(pk=target_pk).first()
    if not source or not target:
        return HttpResponse(tr("Vienas iš sujungiamų įrašų nerastas."), status=404)
    from .permissions import can_see_company, can_see_person

    checker = can_see_person if kind == "person" else can_see_company
    if not (checker(request.user, source) and checker(request.user, target)):
        raise Http404
    if source.merged_into_id == target.pk:
        return redirect(target.get_absolute_url())

    pair_function = all_person_duplicate_pairs if kind == "person" else all_company_duplicate_pairs
    is_duplicate_pair = any(
        {pair["left"].pk, pair["right"].pk} == {source_pk, target_pk}
        for pair in pair_function(duplicate_settings.level)
    )
    if not is_duplicate_pair:
        return HttpResponse(tr("Pasirinkti įrašai pagal dabartines taisykles nėra dublikatai."), status=400)

    from .merging import merge_companies, merge_people
    try:
        target = merge_people(source_pk, target_pk) if kind == "person" else merge_companies(source_pk, target_pk)
    except ValidationError as error:
        messages.error(request, " ".join(error.messages))
        return redirect("contacts:duplicate-list")
    audit_log(AuditLog.MERGE, request=request, target=target,
              old=str(source), new=str(target), detail={"source_id": source_pk, "target_id": target_pk})
    messages.success(request, _("Įrašai sėkmingai sujungti."))
    return redirect(target.get_absolute_url())


@login_required
def contact_detail(request, pk):
    from .detail_editing import grouped_detail_fields
    person = get_object_or_404(visible_people(request.user, Person.objects.select_related("owner", "created_by").prefetch_related("phones", "emails", "addresses", "web_links", "tags", "categories", "activities__created_by", "activities__attachments", "reminders", "company_links__company", "custom_values__field", "responsibles")), pk=pk, deleted_at__isnull=True)
    now = timezone.now()
    open_reminders = person.reminders.filter(completed_at__isnull=True, deleted_at__isnull=True)
    activities = sorted((a for a in person.activities.all() if a.deleted_at is None),
                        key=lambda a: a.created_at, reverse=True)
    return render(request, "contacts/detail.html", {
        "person": person,
        **grouped_detail_fields(person, viewer=request.user),
        "tags": Tag.objects.all(), "categories": Category.objects.all(),
        "reminder_form": ReminderForm(user=request.user),
        "activity_form": ActivityForm(),
        "comment_token": uuid.uuid4().hex,
        "file_token": uuid.uuid4().hex,
        "reminder_token": uuid.uuid4().hex,
        "activity_token": uuid.uuid4().hex,
        "active_reminders": open_reminders,
        "next_reminder": open_reminders.filter(due_at__gt=now).order_by("due_at").first(),
        "overdue_reminder_count": open_reminders.filter(due_at__lte=now).count(),
        "comment_entries": [a for a in activities if a.activity_type == Activity.NOTE],
        "activity_entries": [a for a in activities if a.activity_type != Activity.NOTE],
        "attachment_entries": [att for a in activities for att in a.attachments.all()],
        **_authorship_context(person, "person"),
        **_last_activity_context(activities[0] if activities else None),
    })


@login_required
def contact_type_choice(request):
    return render(request, "contacts/type_choice.html")


@login_required
def contact_create(request):
    form = PersonForm(request.POST or None, user=request.user)
    if request.method == "POST" and form.is_valid():
        duplicate_settings = DuplicateSettings.load()
        duplicates = find_person_duplicates(form.cleaned_data, level=duplicate_settings.level, viewer=request.user) if duplicate_settings.enabled else []
        if duplicates and request.POST.get("confirm_duplicate") != "1":
            return render(request, "contacts/form.html", {"form": form, "title": tr("Pridėti asmenį"), "duplicate_candidates": duplicates})
        person = form.save()
        Person.objects.filter(pk=person.pk).update(created_by=request.user)
        Person.objects.filter(pk=person.pk, owner__isnull=True).update(owner=request.user)
        audit_log(AuditLog.CREATE, request=request, target=person)
        return redirect(person)
    return render(request, "contacts/form.html", {"form": form, "title": tr("Pridėti asmenį")})


@login_required
def contact_edit(request, pk):
    person = get_object_or_404(visible_people(request.user), pk=pk, deleted_at__isnull=True)
    initial = {
        "companies": person.companies.all(),
        "phone": "\n".join(person.phones.values_list("number", flat=True)),
        "email": "\n".join(person.emails.values_list("email", flat=True)),
        "address": "\n".join(person.addresses.values_list("address", flat=True)),
        "url": "\n".join(person.web_links.values_list("url", flat=True)),
    }
    form = PersonForm(request.POST or None, instance=person, initial=initial, user=request.user)
    if request.method == "POST" and form.is_valid():
        duplicate_settings = DuplicateSettings.load()
        duplicates = find_person_duplicates(form.cleaned_data, exclude_pk=person.pk, level=duplicate_settings.level, viewer=request.user) if duplicate_settings.enabled and duplicate_settings.check_on_edit else []
        if duplicates and request.POST.get("confirm_duplicate") != "1":
            return render(request, "contacts/form.html", {"form": form, "title": tr("Redaguoti kontaktą"), "person": person, "duplicate_candidates": duplicates})
        person = form.save()
        for name in form.changed_data:
            audit_log(AuditLog.UPDATE, request=request, target=person, field=name, new=form.cleaned_data.get(name))
        return redirect(person)
    return render(request, "contacts/form.html", {"form": form, "title": tr("Redaguoti kontaktą"), "person": person})


ATTACHMENT_MAX_BYTES = 10 * 1024 * 1024
ATTACHMENT_MAX_FILES = 20            # per upload request
ATTACHMENT_MAX_TOTAL_BYTES = 40 * 1024 * 1024  # per upload request
ATTACHMENT_MAX_PER_RECORD = 200     # across an activity's whole lifetime
ATTACHMENT_ALLOWED_EXTENSIONS = {
    ".jpg", ".jpeg", ".png", ".webp", ".gif", ".pdf", ".txt",
    ".doc", ".docx", ".xls", ".xlsx",
}


def _save_attachments(request, activity):
    """Store the uploaded files that pass the size, type and quota limits.
    Returns the names of any rejected files."""
    rejected = []
    uploads = request.FILES.getlist("attachments")[:ATTACHMENT_MAX_FILES]
    rejected += [u.name for u in request.FILES.getlist("attachments")[ATTACHMENT_MAX_FILES:]]
    total = 0
    existing = Attachment.objects.filter(activity=activity, deleted_at__isnull=True).count()
    for upload in uploads:
        extension = os.path.splitext(upload.name)[1].lower()
        total += upload.size
        if (upload.size > ATTACHMENT_MAX_BYTES or extension not in ATTACHMENT_ALLOWED_EXTENSIONS
                or total > ATTACHMENT_MAX_TOTAL_BYTES or existing >= ATTACHMENT_MAX_PER_RECORD):
            rejected.append(upload.name)
            continue
        Attachment.objects.create(
            activity=activity, file=upload, original_name=upload.name[:255],
            content_type=getattr(upload, "content_type", "") or "", size=upload.size,
        )
        existing += 1
        audit_log(AuditLog.UPDATE, request=request, target=(activity.person or activity.company),
                  field=tr("Priedas"), new=upload.name[:255])
    return rejected


def _report_rejected_attachments(request, rejected):
    if rejected:
        messages.error(request, tr("Nepridėti failai (per dideli arba netinkamo tipo): %(names)s") % {"names": ", ".join(rejected)})


@login_required
def activity_create(request, pk):
    person = get_object_or_404(visible_people(request.user), pk=pk, deleted_at__isnull=True)
    form = ActivityForm(request.POST)
    has_files = bool(request.FILES.getlist("attachments"))
    if form.is_valid() and (form.cleaned_data.get("text") or has_files):
        token = request.POST.get("submission_token", "")
        if token and person.activities.filter(submission_token=token).exists():
            return redirect(person)
        activity = form.save(commit=False)
        activity.person = person
        activity.created_by = request.user
        activity.submission_token = token or None
        activity.text = activity.text or str(tr("Įkeltas failas"))
        activity.save()
        audit_log(AuditLog.CREATE, request=request, target=person,
                  field=str(tr("Veikla")) + ": " + str(activity.get_activity_type_display()), new=activity.text[:200])
        _report_rejected_attachments(request, _save_attachments(request, activity))
    return redirect(person)


@login_required
def activity_edit(request, person_pk, pk):
    person = get_object_or_404(visible_people(request.user), pk=person_pk, deleted_at__isnull=True)
    activity = get_object_or_404(person.activities, pk=pk, deleted_at__isnull=True)
    form = ActivityForm(request.POST or None, instance=activity)
    if request.method == "POST" and form.is_valid():
        old_text = Activity.objects.get(pk=activity.pk).text
        form.save()
        audit_log(AuditLog.UPDATE, request=request, target=person, field=str(tr("Veikla")),
                  old=old_text[:200], new=activity.text[:200])
        messages.success(request, tr("Įrašas atnaujintas."))
        return redirect(person)
    return render(request, "contacts/activity_form.html", {"form": form, "person": person, "activity": activity})


@login_required
def reminder_create(request, pk):
    person = get_object_or_404(visible_people(request.user), pk=pk, deleted_at__isnull=True)
    form = ReminderForm(request.POST, user=request.user)
    if form.is_valid():
        token = request.POST.get("submission_token", "")
        if token and person.reminders.filter(submission_token=token).exists():
            return redirect(person)
        reminder = form.save(commit=False)
        reminder.person = person
        reminder.created_by = request.user
        reminder.assigned_to = reminder.assigned_to or request.user
        reminder.submission_token = token or None
        reminder.save()
        audit_log(AuditLog.CREATE, request=request, target=person, field=str(tr("Priminimas")),
                  new=f"{reminder.text[:150]} · {timezone.localtime(reminder.due_at):%Y-%m-%d %H:%M}")
    return redirect(person)


@login_required
def reminder_complete(request, pk):
    reminder = get_object_or_404(visible_reminders(request.user), pk=pk, deleted_at__isnull=True)
    if request.method == "POST":
        now = timezone.now()
        updated = Reminder.objects.filter(pk=reminder.pk, deleted_at__isnull=True, completed_at__isnull=True).update(
            completed_at=now, updated_at=now,
        )
        if updated:
            audit_log(AuditLog.UPDATE, request=request, target=reminder.record, field=str(tr("Priminimas")),
                      old=reminder.text[:150], new=str(tr("atliktas")))
    return redirect(reminder.record or reverse("contacts:reminder-list"))


@login_required
def reminder_edit(request, pk):
    from .permissions import user_label

    reminder = get_object_or_404(visible_reminders(request.user), pk=pk, deleted_at__isnull=True)
    original_due_at = reminder.due_at
    original_assignee_id = reminder.assigned_to_id
    form = ReminderForm(request.POST or None, instance=reminder, user=request.user)
    if request.method == "POST" and form.is_valid():
        if form.cleaned_data["due_at"] != original_due_at:
            form.instance.read_at = None
        reminder = form.save(commit=False)
        reminder.assigned_to = reminder.assigned_to or request.user
        # Handed to someone else -> it lands unread in their bell.
        handed_over = reminder.assigned_to_id != original_assignee_id and reminder.assigned_to_id != request.user.pk
        if handed_over:
            reminder.read_at = None
        reminder.save()
        if handed_over:
            audit_log(AuditLog.UPDATE, request=request, target=reminder.record, target_type="" if reminder.record else "reminder",
                      target_label="" if reminder.record else reminder.text[:80], field=str(tr("Priskirta")),
                      new=user_label(reminder.assigned_to))
        audit_log(AuditLog.UPDATE, request=request, target=reminder.record, field=str(tr("Priminimas")),
                  old=f"{timezone.localtime(original_due_at):%Y-%m-%d %H:%M}",
                  new=f"{reminder.text[:150]} · {timezone.localtime(reminder.due_at):%Y-%m-%d %H:%M}")
        messages.success(request, tr("Priminimas atnaujintas."))
        return redirect("contacts:reminder-list")
    return render(request, "reminders/form.html", {"form": form, "reminder": reminder})


@login_required
def reminder_delete(request, pk):
    reminder = get_object_or_404(visible_reminders(request.user), pk=pk, deleted_at__isnull=True)
    if request.method == "POST":
        reminder.deleted_at = timezone.now()
        reminder.save(update_fields=["deleted_at", "updated_at"])
        audit_log(AuditLog.DELETE, request=request, target=reminder.record, field=str(tr("Priminimas")),
                  old=reminder.text[:150])
        messages.success(request, tr("Priminimas pašalintas."))
    return redirect("contacts:reminder-list")


REMINDER_SCOPES = ("assigned", "created", "all")


@login_required
def reminder_list(request):
    from .reminder_queries import mine_q, pending_reminders
    now = timezone.now()
    scope = request.GET.get("scope", "assigned")
    if scope not in REMINDER_SCOPES:
        scope = "assigned"
    base = pending_reminders(request.user)
    if scope == "assigned":
        base = base.filter(mine_q(request.user))
    elif scope == "created":
        base = base.filter(created_by=request.user)
    active = base.filter(due_at__lte=now)
    scheduled = base.filter(due_at__gt=now)
    # Opening the list clears the "unread" badge, but only for a genuine same-site
    # visit — a cross-site link or <img> must not silently reset it.
    if request.headers.get("Sec-Fetch-Site", "same-origin") in ("same-origin", "same-site", "none"):
        base.filter(read_at__isnull=True).update(read_at=now)
    return render(request, "reminders/list.html", {
        "active_reminders": active, "scheduled_reminders": scheduled, "reminder_scope": scope,
    })


@login_required
def attachment_download(request, pk):
    available = Attachment.objects.filter(deleted_at__isnull=True, activity__deleted_at__isnull=True).filter(
        Q(activity__person__isnull=False, activity__person__deleted_at__isnull=True)
        | Q(activity__company__isnull=False, activity__company__deleted_at__isnull=True)
    )
    from .permissions import sees_all_records, visible_company_ids, visible_person_ids

    if not sees_all_records(request.user):
        available = available.filter(
            Q(activity__person__pk__in=visible_person_ids(request.user))
            | Q(activity__company__pk__in=visible_company_ids(request.user))
        )
    attachment = get_object_or_404(available, pk=pk)
    if not attachment.file:
        raise Http404
    try:
        return FileResponse(attachment.file.open("rb"), as_attachment=True, filename=attachment.original_name)
    except FileNotFoundError:
        raise Http404("Failas nerastas.")


@login_required
def company_list(request):
    redirect_to = _default_filter_redirect(request, "companies")
    if redirect_to:
        return redirect(redirect_to)
    companies = visible_companies(request.user, Company.objects.filter(deleted_at__isnull=True)).select_related("owner").prefetch_related("people", "tags", "categories")
    filter_values = company_filter_values(request.GET)
    query = filter_values["q"]
    page_size = _list_page_size(request)
    sort_key = request.GET.get("sort", "name")
    direction = "desc" if request.GET.get("direction") == "desc" else "asc"
    companies = apply_company_filters(companies, filter_values, request.user)
    sort_map = {
        "id": "id",
        "name": "name",
        "company_code": "company_code",
        "vat_code": "vat_code",
        "phone": "phone",
        "email": "email",
        "address": "address",
        "contacts": "contact_count",
        "category": "sort_category",
        "tags": "sort_tag",
        "owner": "owner__first_name",
    }
    sort_key = sort_key if sort_key in sort_map else "name"
    order_prefix = "-" if direction == "desc" else ""
    # "Kontaktai" counts only the linked people this user is allowed to see.
    from .permissions import sees_all_records, visible_person_ids

    contact_count_filter = Q(people__deleted_at__isnull=True)
    if not sees_all_records(request.user):
        contact_count_filter &= Q(people__pk__in=visible_person_ids(request.user))
    companies = companies.distinct().annotate(
        contact_count=Count("people", filter=contact_count_filter, distinct=True),
        sort_category=Min("categories__name"),
        sort_tag=Min("tags__name"),
    ).order_by(f"{order_prefix}{sort_map[sort_key]}", "id")
    custom_fields = list(CustomField.objects.filter(entity=CustomField.COMPANY))
    allowed_columns = ["company_code", "vat_code", "address", "phone", "email", "contacts", "owner"] + [field.key for field in custom_fields]
    default_columns = ["company_code", "vat_code", "phone", "email", "contacts"]
    requested_columns = request.GET.getlist("columns")
    if requested_columns:
        columns = [column for column in requested_columns if column in allowed_columns]
        request.session["companies_columns"] = columns
    else:
        columns = request.session.get("companies_columns", default_columns)
        columns = [column for column in columns if column in allowed_columns]
    from django.core.paginator import Paginator

    page = Paginator(companies, page_size).get_page(request.GET.get("page"))
    page.object_list = list(page.object_list)
    custom_columns = [field for field in custom_fields if field.key in columns]
    _attach_custom_cells(page.object_list, custom_columns)
    list_query = request.GET.copy()
    for key in ("page", "sort", "direction"):
        list_query.pop(key, None)
    categories = Category.objects.all()
    tags = Tag.objects.all()
    owner_users = _owner_choices()
    label_maps = {
        "categories": {item.pk: item.name for item in categories},
        "tags": {item.pk: item.name for item in tags},
        "owner": _owner_label_map(owner_users),
        "titles": {"categories": tr("Kategorija"), "tags": tr("Žyma"), **{field.key: field.name for field in custom_fields}},
    }
    return render(request, "companies/list.html", {
        "page": page, "query": query, "page_size": page_size, "sort": sort_key, "direction": direction, "filter_values": filter_values, "columns": columns,
        "custom_fields": custom_fields, "custom_columns": custom_columns,
        "owner_users": owner_users, "can_assign_owner": _can_assign_owner(request.user),
        "owner_filter_hint": _owner_filter_hint(request.user, filter_values["owner"]),
        "page_numbers": _elided_page_numbers(page),
        "list_query": list_query.urlencode(),
        "active_filter_count": active_filter_count(filter_values),
        "filter_chips": filter_chips(request.GET, filter_values, label_maps, request.path),
        "saved_filters": SavedFilter.objects.filter(user=request.user, scope="companies"),
        "tags": tags, "categories": categories,
    })


@login_required
def company_detail(request, pk):
    from .detail_editing import grouped_detail_fields
    company = get_object_or_404(visible_companies(request.user, Company.objects.select_related("owner", "created_by").prefetch_related("person_links__person", "custom_values__field", "responsibles")), pk=pk, deleted_at__isnull=True)
    # Linked contacts, and anything hanging off them, must still respect record visibility.
    visible_linked_ids = list(visible_people(
        request.user, Person.objects.filter(company_links__company=company, deleted_at__isnull=True)
    ).values_list("pk", flat=True))
    linked_people = [link for link in company.person_links.all()
                     if link.person.deleted_at is None and link.person_id in visible_linked_ids]
    history = list(Activity.objects.filter(deleted_at__isnull=True).filter(
        Q(company=company) | Q(person__pk__in=visible_linked_ids)
    ).select_related("person", "company", "created_by").prefetch_related("attachments").distinct().order_by("-created_at"))
    now = timezone.now()
    linked_reminders = Reminder.objects.filter(
        person__pk__in=visible_linked_ids,
        completed_at__isnull=True, deleted_at__isnull=True,
    )
    next_reminder = linked_reminders.filter(due_at__gt=now).select_related("person").order_by("due_at").first()
    return render(request, "companies/detail.html", {
        "company": company,
        **grouped_detail_fields(company, viewer=request.user),
        "tags": Tag.objects.all(), "categories": Category.objects.all(),
        "activity_form": ActivityForm(),
        "comment_token": uuid.uuid4().hex,
        "file_token": uuid.uuid4().hex,
        "activity_token": uuid.uuid4().hex,
        "active_reminders": linked_reminders.select_related("person").order_by("due_at"),
        "next_reminder": next_reminder,
        "overdue_reminder_count": linked_reminders.filter(due_at__lte=now).count(),
        "comment_entries": [a for a in history if a.activity_type == Activity.NOTE],
        "activity_entries": [a for a in history if a.activity_type != Activity.NOTE],
        "attachment_entries": [att for a in history for att in a.attachments.all()],
        "linked_people": linked_people,
        **_authorship_context(company, "company"),
        **_last_activity_context(history[0] if history else None),
    })


@login_required
def company_activity_edit(request, company_pk, pk):
    company = get_object_or_404(visible_companies(request.user), pk=company_pk, deleted_at__isnull=True)
    activity = get_object_or_404(company.activities, pk=pk, deleted_at__isnull=True)
    form = ActivityForm(request.POST or None, instance=activity)
    if request.method == "POST" and form.is_valid():
        old_text = Activity.objects.get(pk=activity.pk).text
        form.save()
        audit_log(AuditLog.UPDATE, request=request, target=company, field=str(tr("Veikla")),
                  old=old_text[:200], new=activity.text[:200])
        messages.success(request, tr("Įrašas atnaujintas."))
        return redirect(company)
    return render(request, "contacts/activity_form.html", {"form": form, "person": company, "activity": activity})


@login_required
def company_activity_create(request, pk):
    company = get_object_or_404(visible_companies(request.user), pk=pk, deleted_at__isnull=True)
    form = ActivityForm(request.POST)
    has_files = bool(request.FILES.getlist("attachments"))
    if form.is_valid() and (form.cleaned_data.get("text") or has_files):
        token = request.POST.get("submission_token", "")
        if token and company.activities.filter(submission_token=token).exists():
            return redirect(company)
        activity = form.save(commit=False)
        activity.company = company
        activity.created_by = request.user
        activity.submission_token = token or None
        activity.text = activity.text or str(tr("Įkeltas failas"))
        activity.save()
        audit_log(AuditLog.CREATE, request=request, target=company,
                  field=str(tr("Veikla")) + ": " + str(activity.get_activity_type_display()), new=activity.text[:200])
        _report_rejected_attachments(request, _save_attachments(request, activity))
    return redirect(company)


@login_required
def company_create(request):
    form = CompanyForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        duplicate_settings = DuplicateSettings.load()
        duplicates = find_company_duplicates(form.cleaned_data, level=duplicate_settings.level, viewer=request.user) if duplicate_settings.enabled else []
        if duplicates and request.POST.get("confirm_duplicate") != "1":
            return render(request, "contacts/form.html", {"form": form, "title": tr("Pridėti įmonę"), "cancel_url": "/companies/", "duplicate_candidates": duplicates})
        company = form.save()
        Company.objects.filter(pk=company.pk).update(created_by=request.user)
        Company.objects.filter(pk=company.pk, owner__isnull=True).update(owner=request.user)
        audit_log(AuditLog.CREATE, request=request, target=company)
        return redirect(company)
    return render(request, "contacts/form.html", {"form": form, "title": tr("Pridėti įmonę"), "cancel_url": "/companies/"})


@login_required
def company_edit(request, pk):
    company = get_object_or_404(visible_companies(request.user), pk=pk, deleted_at__isnull=True)
    form = CompanyForm(request.POST or None, instance=company)
    if request.method == "POST" and form.is_valid():
        duplicate_settings = DuplicateSettings.load()
        duplicates = find_company_duplicates(form.cleaned_data, exclude_pk=company.pk, level=duplicate_settings.level, viewer=request.user) if duplicate_settings.enabled and duplicate_settings.check_on_edit else []
        if duplicates and request.POST.get("confirm_duplicate") != "1":
            return render(request, "contacts/form.html", {"form": form, "title": tr("Redaguoti įmonę"), "company": company, "cancel_url": company.get_absolute_url(), "duplicate_candidates": duplicates})
        company = form.save()
        for name in form.changed_data:
            audit_log(AuditLog.UPDATE, request=request, target=company, field=name, new=form.cleaned_data.get(name))
        return redirect(company)
    return render(request, "contacts/form.html", {"form": form, "title": tr("Redaguoti įmonę"), "company": company, "cancel_url": company.get_absolute_url()})


def _value(row, *names):
    for name in names:
        value = row.get(name)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def _import_relation_names(row, *names):
    values = []
    for value in _value(row, *names).split(";"):
        value = value.strip()
        if value and value not in values:
            values.append(value)
    return values


@login_required
def contacts_export(request):
    _require_capability(request, "can_export")
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="crm-kontaktai.csv"'
    response.write("\ufeff")
    writer = csv.writer(response)
    writer.writerow(["Vardas", "Pavardė", "Pareigos", "Įmonė", "Telefonai", "El. paštai", "Adresai", "URL", "Tagai", "Kategorijos", "Atsakingas", "Atsakingi"])
    people = visible_people(request.user, Person.objects.filter(deleted_at__isnull=True))
    if request.method == "POST":
        people = people.filter(pk__in=request.POST.getlist("selected"))
    people = people.select_related("owner").prefetch_related("phones", "emails", "addresses", "web_links", "tags", "categories", "company_links__company", "responsibles")
    from .permissions import user_label
    rows = 0
    for person in people:
        writer.writerow([csv_safe(v) for v in [person.first_name, person.last_name, person.job_title, "; ".join(link.company.name for link in person.company_links.all()), "; ".join(item.number for item in person.phones.all()), "; ".join(item.email for item in person.emails.all()), "; ".join(item.address for item in person.addresses.all()), "; ".join(item.url for item in person.web_links.all()), "; ".join(item.name for item in person.tags.all()), "; ".join(item.name for item in person.categories.all()), user_label(person.owner), "; ".join(u.get_username() for u in person.responsibles.all())]])
        rows += 1
    audit_log(AuditLog.EXPORT, request=request, target_type="export", target_label=str(tr("Kontaktai (CSV)")), new=str(rows))
    return response

@login_required
def companies_export(request):
    _require_capability(request, "can_export")
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="crm-imones.csv"'
    response.write("\ufeff")
    writer = csv.writer(response)
    writer.writerow(["Pavadinimas", "Įmonės kodas", "PVM kodas", "Adresas", "Telefonas", "El. paštas", "Atsakingas", "Atsakingi"])
    companies = visible_companies(request.user, Company.objects.filter(deleted_at__isnull=True)).select_related("owner").prefetch_related("responsibles")
    if request.method == "POST": companies = companies.filter(pk__in=request.POST.getlist("selected"))
    from .permissions import user_label
    rows = 0
    for company in companies:
        writer.writerow([csv_safe(v) for v in [company.name, company.company_code, company.vat_code, company.address, company.phone, company.email, user_label(company.owner), "; ".join(u.get_username() for u in company.responsibles.all())]])
        rows += 1
    audit_log(AuditLog.EXPORT, request=request, target_type="export", target_label=str(tr("Įmonės (CSV)")), new=str(rows))
    return response


def _resolve_import_owner(value, cache):
    """Match an "Atsakingas" cell to an active user by username, full name or email."""
    value = (value or "").strip()
    if not value:
        return None
    if value in cache:
        return cache[value]
    users = get_user_model().objects.filter(is_active=True)
    user = (
        users.filter(username__iexact=value).first()
        or users.filter(email__iexact=value).first()
        or next((u for u in users if u.get_full_name().strip().lower() == value.lower()), None)
    )
    cache[value] = user
    return user


# Canonical import field -> (human label, accepted header aliases). The first
# alias is the key that _value()/_import_relation_names() look for first, so the
# column-mapping step rewrites every row to use it.
IMPORT_COLUMNS = [
    ("Vardas", tr("Vardas"), ("Vardas", "first_name", "First name")),
    ("Pavardė", tr("Pavardė"), ("Pavardė", "last_name", "Last name")),
    ("Pareigos", tr("Pareigos"), ("Pareigos", "job_title", "Title")),
    ("Įmonė", tr("Įmonė"), ("Įmonė", "company", "Company")),
    ("Telefonai", tr("Telefonai"), ("Telefonai", "Telefonas", "phone", "Phone")),
    ("El. paštai", tr("El. paštai"), ("El. paštai", "El. paštas", "email", "Email")),
    ("Adresai", tr("Adresai"), ("Adresai", "Adresas", "address", "Address")),
    ("URL", tr("Nuorodos"), ("URL", "url", "Website")),
    ("Tagai", tr("Žymos"), ("Tagai", "Žymos", "Tags", "tags")),
    ("Kategorijos", tr("Kategorijos"), ("Kategorijos", "Categories", "categories")),
    ("Atsakingas", tr("Atsakingas (pagrindinis)"), ("Atsakingas", "owner", "Owner")),
    ("Atsakingi", tr("Atsakingi (papildomi)"), ("Atsakingi", "responsibles", "Responsibles")),
]
IMPORT_FIELD_KEYS = [key for key, _label, _aliases in IMPORT_COLUMNS]


def _guess_import_mapping(headers):
    """Best-effort {csv header: canonical field} from header names."""
    lookup = {}
    for key, _label, aliases in IMPORT_COLUMNS:
        for alias in aliases:
            lookup[alias.strip().lower()] = key
    mapping = {}
    for header in headers:
        match = lookup.get(str(header).strip().lower())
        if match and match not in mapping.values():
            mapping[header] = match
    return mapping


def _apply_import_mapping(rows, mapping):
    """Rewrite each row so its keys are canonical field names (per `mapping`)."""
    if not mapping:
        return rows
    remapped = []
    for row in rows:
        new_row = {}
        for header, field in mapping.items():
            if field and header in row:
                new_row[field] = row[header]
        remapped.append(new_row)
    return remapped


def _import_one_row(row, owner, mode, owner_cache, can_reassign=True):
    """Import a single already-mapped row. Returns 'created' / 'updated' / 'skipped'."""
    first_name = _value(row, "Vardas", "first_name", "First name")
    last_name = _value(row, "Pavardė", "last_name", "Last name")
    if not first_name and not last_name:
        return "skipped", None
    email_values = list(filter(None, (item.strip() for item in _value(row, "El. paštai", "El. paštas", "email", "Email").split(";"))))
    email = email_values[0] if email_values else ""
    person = None
    if mode != "new":
        from .permissions import visible_people

        base = Person.objects.filter(deleted_at__isnull=True)
        match = base.filter(emails__email__iexact=email) if email else base.none()
        if not match.exists():
            match = base.filter(first_name__iexact=first_name, last_name__iexact=last_name)
        person = visible_people(owner, match).first()
        # A row that matches a record the importer may not see must not silently
        # update it (and possibly reassign its owner) — that would expose the card.
        if person is None and match.exists():
            raise ValueError(str(tr("Kontaktas su šiuo el. paštu arba vardu jau yra, bet jums nematomas.")))
    if person and mode == "skip":
        return "skipped", None
    tag_names = _import_relation_names(row, "Tagai", "Žymos", "Tags", "tags")
    category_names = _import_relation_names(row, "Kategorijos", "Categories", "categories")
    if len(tag_names) > 3 or len(category_names) > 3:
        raise ValueError(str(tr("Viršytas leistinas žymų arba kategorijų skaičius (daugiausia 3)")))
    values = {"first_name": first_name, "last_name": last_name, "job_title": _value(row, "Pareigos", "job_title")}
    # Owner reassignment via import follows the same capability as the inline editor.
    row_owner = _resolve_import_owner(_value(row, "Atsakingas", "owner", "Owner"), owner_cache) if can_reassign else None
    if person:
        for field, value in values.items():
            if value:
                setattr(person, field, value)
        if row_owner:
            person.owner = row_owner
        person.save()
        outcome = "updated"
    else:
        person = Person.objects.create(owner=row_owner or owner, created_by=owner, **values)
        outcome = "created"
    for company_name in _import_relation_names(row, "Įmonė", "company", "Company"):
        company, _ = Company.objects.get_or_create(name=company_name)
        PersonCompanyLink.objects.get_or_create(
            person=person, company=company,
            defaults={"is_primary": not person.company_links.filter(is_primary=True).exists()},
        )
    for number in filter(None, (item.strip() for item in _value(row, "Telefonai", "Telefonas", "phone", "Phone").split(";"))):
        PhoneNumber.objects.get_or_create(person=person, number=number, defaults={"is_primary": not person.phones.exists()})
    for address in filter(None, (item.strip() for item in _value(row, "Adresai", "Adresas", "address", "Address").split(";"))):
        PostalAddress.objects.get_or_create(person=person, address=address)
    for url in filter(None, (safe_url(item.strip()) for item in _value(row, "URL", "url", "Website").split(";"))):
        WebLink.objects.get_or_create(person=person, url=url)
    for address in email_values:
        EmailAddress.objects.get_or_create(person=person, email=address, defaults={"is_primary": not person.emails.exists()})
    for tag_name in tag_names:
        tag, _ = Tag.objects.get_or_create(name=tag_name[:60])
        person.tags.add(tag)
    for category_name in category_names:
        category, _ = Category.objects.get_or_create(name=category_name[:60])
        person.categories.add(category)
    if can_reassign:
        extra = [_resolve_import_owner(name, owner_cache) for name in _import_relation_names(row, "Atsakingi", "responsibles", "Responsibles")]
        extra = [u for u in extra if u and u.pk != person.owner_id]
        if extra:
            person.responsibles.add(*extra)
    return outcome, person


@transaction.atomic
def _import_contact_rows(rows, owner=None, *, mode="update", collect_errors=False):
    """Import mapped rows. With collect_errors=True a failing row is recorded and
    skipped instead of aborting the whole batch."""
    created = updated = skipped = possible_duplicates = 0
    created_person_ids = set()
    owner_cache = {}
    errors = []
    duplicate_settings = DuplicateSettings.load()
    report_duplicates = duplicate_settings.enabled and duplicate_settings.check_on_import
    from .permissions import has_capability

    can_reassign = owner is None or has_capability(owner, "can_reassign_owner")
    for index, row in enumerate(rows, start=2):  # row 1 is the header line
        try:
            with transaction.atomic():
                outcome, person = _import_one_row(row, owner, mode, owner_cache, can_reassign)
        except Exception as error:
            if not collect_errors:
                raise
            errors.append({"row": index, "error": str(error), "data": row})
            continue
        if outcome == "created":
            created += 1
            created_person_ids.add(person.pk)
        elif outcome == "updated":
            updated += 1
        else:
            skipped += 1
    if report_duplicates and created_person_ids:
        possible_duplicates = sum(
            1 for pair in all_person_duplicate_pairs(duplicate_settings.level)
            if pair["left"].pk in created_person_ids or pair["right"].pk in created_person_ids
        )
    return {"created": created, "updated": updated, "skipped": skipped, "possible_duplicates": possible_duplicates,
            "duplicate_check_enabled": report_duplicates, "errors": errors, "error_count": len(errors)}


IMPORT_PREVIEW_LIMIT = 5000


def _decode_csv(data, setting):
    """Decode CSV bytes using the configured encoding, or try common ones."""
    candidates = ["utf-8-sig", "cp1257"] if setting == "auto" else [setting]
    for encoding in candidates:
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace")


def _csv_delimiter(text, setting):
    """Return the configured CSV delimiter, or sniff one from the header line."""
    if setting == "tab":
        return "\t"
    if setting in (",", ";"):
        return setting
    sample = text[:4096]
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
    except csv.Error:
        return ";" if sample.count(";") > sample.count(",") else ","


def _read_import_rows(upload):
    name = upload.name.lower()
    if name.endswith(".csv"):
        system = SystemSettings.load()
        text = _decode_csv(upload.file.read(), system.import_encoding)
        delimiter = _csv_delimiter(text, system.import_delimiter)
        raw = list(csv.DictReader(io.StringIO(text), delimiter=delimiter))
    elif name.endswith(".xlsx"):
        from openpyxl import load_workbook

        sheet = load_workbook(upload, read_only=True, data_only=True).active
        headers = [str(cell.value or "").strip() for cell in next(sheet.iter_rows())]
        raw = [{headers[index]: cell.value for index, cell in enumerate(row)} for row in sheet.iter_rows(min_row=2)]
    else:
        raise ValueError(tr("Netinkamas failo formatas"))
    # Normalise every value to a string so the rows are JSON/session safe.
    return [{str(key): ("" if value is None else str(value)) for key, value in row.items() if key} for row in raw]


def _import_preview(rows, mapping):
    mapped = _apply_import_mapping(rows, mapping)
    created = updated = skipped = problems = 0
    for row in mapped:
        first_name = _value(row, "Vardas", "first_name", "First name")
        last_name = _value(row, "Pavardė", "last_name", "Last name")
        emails = list(filter(None, (item.strip() for item in _value(row, "El. paštai", "El. paštas", "email", "Email").split(";"))))
        if not first_name and not last_name:
            skipped += 1
            continue
        if len(_import_relation_names(row, "Tagai", "Žymos", "Tags", "tags")) > 3 or len(_import_relation_names(row, "Kategorijos", "Categories", "categories")) > 3:
            problems += 1
        match = (emails and Person.objects.filter(emails__email__iexact=emails[0], deleted_at__isnull=True).exists()) or \
            Person.objects.filter(first_name__iexact=first_name, last_name__iexact=last_name, deleted_at__isnull=True).exists()
        if match:
            updated += 1
        else:
            created += 1
    headers = list(rows[0].keys()) if rows else []
    return {"total": len(rows), "created": created, "updated": updated, "skipped": skipped, "problems": problems,
            "headers": headers, "sample": [[row.get(header, "") for header in headers] for row in rows[:8]],
            "columns": [{"key": key, "label": label} for key, label, _aliases in IMPORT_COLUMNS],
            "mapping": mapping}


@login_required
def contacts_import(request):
    _require_capability(request, "can_import")
    context = {}
    if request.method == "POST" and request.POST.get("confirm") == "1":
        rows = request.session.pop("import_rows", None)
        stored_mapping = request.session.pop("import_mapping", {})
        if not rows:
            messages.error(request, tr("Importo peržiūra pasibaigė. Įkelkite failą iš naujo."))
        else:
            headers = list(rows[0].keys()) if rows else []
            submitted = {header: request.POST.get("map_" + header, "").strip() for header in headers if request.POST.get("map_" + header, "").strip() in IMPORT_FIELD_KEYS}
            mapping = submitted or stored_mapping
            mode = request.POST.get("dedup") if request.POST.get("dedup") in {"skip", "update", "new"} else "update"
            try:
                result = _import_contact_rows(_apply_import_mapping(rows, mapping), owner=request.user, mode=mode, collect_errors=True)
            except Exception:
                messages.error(request, tr("Nepavyko importuoti. Patikrinkite stulpelius ir failo formatą."))
            else:
                context["result"] = result
                if result["errors"]:
                    request.session["import_errors"] = [{"row": e["row"], "error": e["error"], "data": e["data"]} for e in result["errors"]]
                else:
                    request.session.pop("import_errors", None)
                audit_log(AuditLog.IMPORT, request=request, target_type="import", target_label=str(tr("Kontaktų importas")),
                          new=str(tr("sukurta %(c)s, atnaujinta %(u)s, praleista %(s)s, klaidų %(e)s")) % {"c": result["created"], "u": result["updated"], "s": result["skipped"], "e": result["error_count"]},
                          detail={k: v for k, v in result.items() if k != "errors"})
    elif request.method == "POST":
        upload = request.FILES.get("file")
        if not upload or upload.size > 10 * 1024 * 1024:
            messages.error(request, tr("Pasirinkite iki 10 MB dydžio CSV arba XLSX failą."))
        else:
            try:
                rows = _read_import_rows(upload)
            except Exception:
                rows = None
                messages.error(request, tr("Nepavyko perskaityti failo. Patikrinkite stulpelius ir failo formatą."))
            if rows is not None and len(rows) > IMPORT_PREVIEW_LIMIT:
                messages.error(request, tr("Peržiūrai skirtas failas su ne daugiau kaip %(n)s eilučių.") % {"n": IMPORT_PREVIEW_LIMIT})
            elif rows:
                headers = list(rows[0].keys())
                mapping = _guess_import_mapping(headers)
                request.session["import_rows"] = rows
                request.session["import_mapping"] = mapping
                context["preview"] = _import_preview(rows, mapping)
            elif rows is not None:
                messages.error(request, tr("Faile nerasta įrašų."))
    else:
        request.session.pop("import_rows", None)
        request.session.pop("import_mapping", None)
    context["has_error_report"] = bool(request.session.get("import_errors"))
    return render(request, "import_export.html", context)


@login_required
def contacts_import_errors(request):
    _require_capability(request, "can_import")
    errors = request.session.get("import_errors")
    if not errors:
        raise Http404
    field_names = []
    for entry in errors:
        for key in entry["data"].keys():
            if key not in field_names:
                field_names.append(key)
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="importo-klaidos.csv"'
    response.write("﻿")
    writer = csv.writer(response)
    writer.writerow([str(tr("Eilutė")), str(tr("Klaida"))] + field_names)
    for entry in errors:
        writer.writerow([entry["row"], csv_safe(entry["error"])] + [csv_safe(entry["data"].get(name, "")) for name in field_names])
    return response
